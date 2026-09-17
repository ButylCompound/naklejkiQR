import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import gui
from reconciliation import reconcile_raw_materials


class RawMaterialsMatchingTests(unittest.TestCase):
    def _write_expected(self, path, items):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Stan magazynowy surowców'
        sheet['A3'] = 'Row Labels'
        sheet['B3'] = 'Quantity'
        for row_number, (name, quantity) in enumerate(items, 4):
            sheet.cell(row=row_number, column=1, value=name).alignment = Alignment(indent=1)
            sheet.cell(row=row_number, column=2, value=quantity)
        workbook.save(path)

    def _write_scans(self, path, items):
        pd.DataFrame(
            [
                {'Lp': index, 'Produkt': name, 'Ilość': quantity}
                for index, (name, quantity) in enumerate(items, 1)
            ]
        ).to_csv(path, sep=';', index=False)

    def test_exact_name_wins_over_longer_fuzzy_candidate(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'
            short_name = 'P1M22 135/3,0/1100'
            long_name = short_name + ' poza tolerancją'
            self._write_expected(excel_path, [(short_name, 0), (long_name, 100)])
            self._write_scans(csv_path, [(short_name, 25)])

            result = reconcile_raw_materials(excel_path, csv_path)
            records = {record['Produkt']: record for record in result['records']}

            self.assertEqual(records[short_name]['Zeskanowana Ilość'], 25)
            self.assertEqual(records[long_name]['Zeskanowana Ilość'], 0)

    def test_polish_only_name_does_not_crash_or_become_empty(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'
            self._write_expected(excel_path, [('ABC', 10), ('ĄĘĆ', 5)])
            self._write_scans(csv_path, [('ĄĘĆ', 5)])

            result = reconcile_raw_materials(excel_path, csv_path)
            records = {record['Produkt']: record for record in result['records']}

            self.assertEqual(records['ĄĘĆ']['Zeskanowana Ilość'], 5)
            self.assertEqual(records['ĄĘĆ']['Kategoria'], 'zgodne')

    def test_numbered_pack_labels_aggregate_to_the_base_material(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'
            self._write_expected(excel_path, [('Raw Material', 30)])
            self._write_scans(
                csv_path,
                [('Raw Material (1)', 10), ('Raw Material (2)', 20)],
            )

            result = reconcile_raw_materials(excel_path, csv_path)
            record = result['records'][0]

            self.assertEqual(record['Produkt'], 'Raw Material')
            self.assertEqual(record['Zeskanowana Ilość'], 30)
            self.assertEqual(record['Kategoria'], 'zgodne')

    def test_totals_statuses_duplicate_expected_rows_and_footer(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'
            self._write_expected(excel_path, [
                ('Matched', 40),
                ('Matched', 60),
                ('Missing', 50),
                ('Surplus', 25),
            ])
            pd.DataFrame([
                {'Lp': 1, 'Produkt': 'Matched', 'Ilość': 100},
                {'Lp': 2, 'Produkt': 'Missing', 'Ilość': 20},
                {'Lp': 3, 'Produkt': 'Surplus', 'Ilość': 40},
                {'Lp': 4, 'Produkt': 'Unknown', 'Ilość': 5},
                {'Lp': 'Liczba palet', 'Produkt': 'Footer', 'Ilość': 'invalid'},
            ]).to_csv(csv_path, sep=';', index=False)

            result = reconcile_raw_materials(excel_path, csv_path)
            records = {record['Produkt']: record for record in result['records']}

            self.assertEqual(result['summary'], {
                'total_expected': 175,
                'total_scanned': 165,
                'matched_groups': 1,
                'missing_groups': 1,
                'surplus_groups': 2,
            })
            self.assertEqual(records['Matched']['Kategoria'], 'zgodne')
            self.assertEqual(records['Missing']['Różnica'], -30)
            self.assertEqual(records['Surplus']['Różnica'], 15)
            self.assertEqual(records['Unknown (Brak w bazie)']['Zeskanowana Ilość'], 5)
            self.assertEqual([record['Lp'] for record in result['records']], [1, 2, 3, 4])

    def test_alternate_sheet_and_quantity_column_are_supported(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'custom'
            sheet['A3'] = 'Row Labels'
            sheet['B3'] = 'Quantity'
            sheet['A4'] = 'Materiał Łódź'
            sheet['A4'].alignment = Alignment(indent=1)
            sheet['B4'] = 10.5
            workbook.save(excel_path)
            pd.DataFrame([
                {'Lp': 1, 'Produkt': 'Material Lodz', 'Waga': '10,5 kg'},
            ]).to_csv(csv_path, sep=';', index=False)

            result = reconcile_raw_materials(excel_path, csv_path, excel_sheet='custom')

            self.assertEqual(result['summary']['matched_groups'], 1)
            self.assertEqual(result['records'][0]['Zeskanowana Ilość'], 10.5)

    def test_difference_below_one_hundredth_is_treated_as_equal(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            excel_path = temp_path / 'expected.xlsx'
            csv_path = temp_path / 'scans.csv'
            self._write_expected(excel_path, [('Material', 10)])
            self._write_scans(csv_path, [('Material', 10.005)])

            result = reconcile_raw_materials(excel_path, csv_path)

            self.assertEqual(result['summary']['matched_groups'], 1)
            self.assertEqual(result['records'][0]['Kategoria'], 'zgodne')


class ExcelUpdateTests(unittest.TestCase):
    def test_confirmation_uses_matching_records(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx') as excel_file:
            fake_window = SimpleNamespace(
                tab_palety=SimpleNamespace(
                    all_records=[{'Kategoria': 'zgodne', 'Excel_Row': 2}]
                ),
                entry_excel=SimpleNamespace(get=lambda: excel_file.name),
            )

            with patch.object(gui.messagebox, 'askyesno', return_value=False) as confirm:
                gui.MainWindow.update_excel_dates(fake_window)

            self.assertIn('dla 1 zgodnych palet', confirm.call_args.args[1])


if __name__ == '__main__':
    unittest.main()
