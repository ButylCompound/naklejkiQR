import sys
import tempfile
import unittest
from datetime import date, time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from reconciliation import reconcile_inventory, reconcile_raw_materials


class InputValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.excel_path = self.root / 'input.xlsx'
        self.csv_path = self.root / 'input.csv'

    def tearDown(self):
        self.temp_dir.cleanup()

    def write_pallet_excel(self, **changes):
        row = {
            'Nazwa handlowa': 'Product',
            'ilość na palecie': 10,
            'Stan': 'Stan',
            'Data': date(2026, 8, 25),
            'Godzina końca palety': time(9, 5),
        }
        row.update(changes)
        pd.DataFrame([row]).to_excel(self.excel_path, sheet_name='palety', index=False)

    def write_pallet_csv(self, **changes):
        row = {
            'Lp': 1,
            'Produkt': 'Product',
            'Ilość': 10,
            'Data naklejki': '2026-08-25 09:05:00',
        }
        row.update(changes)
        pd.DataFrame([row]).to_csv(self.csv_path, sep=';', index=False)

    def write_raw_excel(self, items=(('Material', 10),), sheet='Stan magazynowy surowców'):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet
        worksheet['A3'] = 'Row Labels'
        worksheet['B3'] = 'Quantity'
        for row_number, (name, quantity) in enumerate(items, 4):
            worksheet.cell(row=row_number, column=1, value=name).alignment = Alignment(indent=1)
            worksheet.cell(row=row_number, column=2, value=quantity)
        workbook.save(self.excel_path)

    def write_raw_csv(self, **changes):
        row = {'Lp': 1, 'Produkt': 'Material', 'Ilość': 10}
        row.update(changes)
        pd.DataFrame([row]).to_csv(self.csv_path, sep=';', index=False)

    def assert_runtime_error(self, function, text):
        with self.assertRaisesRegex(RuntimeError, text):
            function(self.excel_path, self.csv_path)

    def test_missing_or_corrupt_pallet_excel_is_rejected(self):
        self.write_pallet_csv()

        with self.assertRaisesRegex(RuntimeError, 'pliku Excel'):
            reconcile_inventory(self.root / 'missing.xlsx', self.csv_path)

        self.excel_path.write_text('not an Excel file', encoding='utf-8')
        self.assert_runtime_error(reconcile_inventory, 'pliku Excel')

    def test_pallet_excel_without_palety_sheet_is_rejected(self):
        pd.DataFrame({'A': [1]}).to_excel(self.excel_path, sheet_name='wrong', index=False)
        self.write_pallet_csv()

        self.assert_runtime_error(reconcile_inventory, 'pliku Excel')

    def test_pallet_excel_missing_each_required_column_is_rejected(self):
        required = (
            'Nazwa handlowa',
            'ilość na palecie',
            'Stan',
            'Data',
            'Godzina końca palety',
        )
        self.write_pallet_csv()
        for column in required:
            with self.subTest(column=column):
                self.write_pallet_excel()
                dataframe = pd.read_excel(self.excel_path, sheet_name='palety')
                dataframe.drop(columns=[column]).to_excel(self.excel_path, sheet_name='palety', index=False)
                self.assert_runtime_error(reconcile_inventory, column)

    def test_pallet_excel_with_invalid_quantity_date_or_time_is_rejected(self):
        self.write_pallet_csv()
        cases = (
            ({'ilość na palecie': 'heavy'}, 'Nieprawidłowa ilość'),
            ({'Data': 'not-a-date'}, 'Nieprawidłowa data'),
            ({'Godzina końca palety': 'not-a-time'}, 'Nieprawidłowa godzina'),
        )
        for changes, message in cases:
            with self.subTest(changes=changes):
                self.write_pallet_excel(**changes)
                self.assert_runtime_error(reconcile_inventory, message)

    def test_missing_empty_or_wrongly_delimited_pallet_csv_is_rejected(self):
        self.write_pallet_excel()

        with self.assertRaisesRegex(RuntimeError, 'pliku CSV'):
            reconcile_inventory(self.excel_path, self.root / 'missing.csv')

        self.csv_path.write_text('', encoding='utf-8')
        self.assert_runtime_error(reconcile_inventory, 'pliku CSV')

        self.csv_path.write_text(
            'Lp,Produkt,Ilość,Data naklejki\n1,Product,10,2026-08-25 09:05:00\n',
            encoding='utf-8',
        )
        self.assert_runtime_error(reconcile_inventory, 'wymaganych kolumn')

    def test_pallet_csv_missing_required_columns_is_rejected(self):
        self.write_pallet_excel()
        for column in ('Produkt', 'Ilość', 'Data naklejki'):
            with self.subTest(column=column):
                self.write_pallet_csv()
                dataframe = pd.read_csv(self.csv_path, sep=';').drop(columns=[column])
                dataframe.to_csv(self.csv_path, sep=';', index=False)
                expected = 'kolumny z ilością' if column == 'Ilość' else column
                self.assert_runtime_error(reconcile_inventory, expected)

    def test_pallet_csv_with_invalid_quantity_or_date_is_rejected(self):
        self.write_pallet_excel()
        for changes, message in (
            ({'Ilość': 'many'}, 'Nieprawidłowa ilość'),
            ({'Data naklejki': 'yesterday'}, 'Nieprawidłowa data naklejki'),
        ):
            with self.subTest(changes=changes):
                self.write_pallet_csv(**changes)
                self.assert_runtime_error(reconcile_inventory, message)

    def test_missing_corrupt_or_wrong_sheet_raw_excel_is_rejected(self):
        self.write_raw_csv()

        with self.assertRaisesRegex(RuntimeError, 'pliku Excel'):
            reconcile_raw_materials(self.root / 'missing.xlsx', self.csv_path)

        self.excel_path.write_text('not an Excel file', encoding='utf-8')
        self.assert_runtime_error(reconcile_raw_materials, 'pliku Excel')

        self.write_raw_excel(sheet='wrong')
        self.assert_runtime_error(reconcile_raw_materials, 'pliku Excel')

    def test_raw_excel_with_wrong_layout_or_invalid_quantity_is_rejected(self):
        self.write_raw_csv()

        self.write_raw_excel(items=())
        self.assert_runtime_error(reconcile_raw_materials, 'oczekiwanym układzie')

        self.write_raw_excel(items=(('Material', 'many'),))
        self.assert_runtime_error(reconcile_raw_materials, 'Nieprawidłowa ilość')

    def test_raw_csv_missing_columns_wrong_delimiter_or_invalid_quantity_is_rejected(self):
        self.write_raw_excel()

        with self.assertRaisesRegex(RuntimeError, 'pliku CSV'):
            reconcile_raw_materials(self.excel_path, self.root / 'missing.csv')

        self.csv_path.write_text('', encoding='utf-8')
        self.assert_runtime_error(reconcile_raw_materials, 'pliku CSV')

        for column, message in (('Produkt', 'Produkt'), ('Ilość', 'kolumny z ilością')):
            with self.subTest(column=column):
                self.write_raw_csv()
                dataframe = pd.read_csv(self.csv_path, sep=';').drop(columns=[column])
                dataframe.to_csv(self.csv_path, sep=';', index=False)
                self.assert_runtime_error(reconcile_raw_materials, message)

        self.csv_path.write_text('Lp,Produkt,Ilość\n1,Material,10\n', encoding='utf-8')
        self.assert_runtime_error(reconcile_raw_materials, 'Produkt')

        self.write_raw_csv(Ilość='many')
        self.assert_runtime_error(reconcile_raw_materials, 'Nieprawidłowa ilość')


if __name__ == '__main__':
    unittest.main()
